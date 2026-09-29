import asyncio
from io import BytesIO

import fitz
import pytest
from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.api import batches, exports, scans
from app.core.exceptions import ScannerError
from app.db.base import Base
from app.models import ScanAttempt, ScanBatch, ScanItem
from app.schemas.kk import KKUpdate
from app.services.pdf_extraction_service import extract_pdf_document
from app.services.validation_service import validate_extraction


def _insert_centered(page, center_x, y, text, fontsize=6):
    width = fitz.get_text_length(text, fontsize=fontsize)
    page.insert_text((center_x - width / 2, y), text, fontsize=fontsize)


def synthetic_kk_pdf(
    pages=1, include_text=True, missing_secondary_row=False,
    include_second_member=True, primary_edges=None, secondary_anchor_y=294,
    first_marital_status='KAWIN TERCATAT', first_marriage_date='01-01-2000',
    secondary_edges=None, draw_grid=False,
):
    """Fictitious residents on the same landscape grid as the supplied blanko."""
    document = fitz.open()
    for _ in range(pages):
        page = document.new_page(width=841.9, height=595.3)
        if not include_text:
            continue
        page_width = 841.9
        primary_edges = primary_edges or [25, 202, 282, 324, 425, 467, 518, 640, 783, page_width]
        secondary_edges = secondary_edges or [25, 109, 160, 257, 337, 402, 467, 640, page_width]
        header = [
            (280, 35, 'KARTU KELUARGA'), (250, 57, 'NO : 3311111111111111'),
            (120, 70, 'Nama Kepala Keluarga'), (227, 70, ': BUDI SANTOSO'),
            (120, 80, 'Alamat'), (227, 80, ': JL MELATI 1'),
            (120, 90, 'RT/RW'), (227, 90, ': 001 / 002'),
            (120, 100, 'Kode Pos'), (227, 100, ': 59511'),
            (505, 70, 'Desa/Kelurahan'), (579, 70, ': LAMBANGGELUN'),
            (505, 80, 'Kecamatan'), (579, 80, ': KARANGANYAR'),
            (505, 90, 'Kabupaten/Kota'), (579, 90, ': DEMAK'),
            (505, 100, 'Provinsi'), (579, 100, ': JAWA TENGAH'),
            (83, 121, 'NAMA LENGKAP'), (234, 121, 'NIK'),
            (260, 277, 'KEWARGANEGARAAN'),
        ]
        for x, y, text in header:
            page.insert_text((x, y), text, fontsize=6)
        for number, (left, right) in enumerate(zip(primary_edges, primary_edges[1:]), start=1):
            _insert_centered(page, (left + right) / 2, 137, f'({number})')
        for number, (left, right) in enumerate(zip(secondary_edges, secondary_edges[1:]), start=10):
            _insert_centered(page, (left + right) / 2, secondary_anchor_y, f'({number})')
        if draw_grid:
            for x in [10, *primary_edges]:
                page.draw_line(fitz.Point(x, 112), fitz.Point(x, 260), width=0.5)
            for x in [10, *secondary_edges]:
                page.draw_line(
                    fitz.Point(x, secondary_anchor_y - 45),
                    fitz.Point(x, secondary_anchor_y + 125), width=0.5,
                )

        primary_values = [
            'BUDI SANTOSO', '3311111111111112', 'LAKI-LAKI', 'DEMAK',
            '01-01-1980', 'ISLAM', 'SMA', 'PETANI', 'O',
        ]
        secondary_values = [
            first_marital_status, first_marriage_date, 'KEPALA KELUARGA',
            'WNI', '-', '-', 'SUGENG', 'SRI',
        ]
        page.insert_text((16, 153), '1', fontsize=6)
        for left, text in zip(primary_edges, primary_values):
            page.insert_text((left + 3, 153), text, fontsize=6)
        page.insert_text((16, secondary_anchor_y + 16), '1', fontsize=6)
        for left, text in zip(secondary_edges, secondary_values):
            if text:
                page.insert_text((left + 3, secondary_anchor_y + 16), text, fontsize=6)
        if include_second_member:
            page.insert_text((16, 164), '2', fontsize=6)
            values = ['ANI SANTOSO', '3311111111111113', 'PEREMPUAN', 'DEMAK', '02-02-2005', 'ISLAM', 'SMA', 'PELAJAR', 'A']
            for left, text in zip(primary_edges, values):
                page.insert_text((left + 3, 164), text, fontsize=6)
        if include_second_member and not missing_secondary_row:
            page.insert_text((16, secondary_anchor_y + 27), '2', fontsize=6)
            values = ['BELUM KAWIN', '-', 'ANAK', 'WNI', '-', '-', 'BUDI SANTOSO', 'SRI']
            for left, text in zip(secondary_edges, values):
                page.insert_text((left + 3, secondary_anchor_y + 27), text, fontsize=6)
        # Preprinted empty slot must not become a resident.
        for y in (175, secondary_anchor_y + 38):
            page.insert_text((16, y), '3', fontsize=6)
            edges = primary_edges if y == 175 else secondary_edges
            page.insert_text((edges[0] + 3, y), '-', fontsize=6)
            page.insert_text((edges[1] + 3, y), '-', fontsize=6)
        # Even a text watermark is ignored; the supplied PDF uses an image.
        page.insert_text((400, 480), 'DRAFT', fontsize=20)
    data = document.tobytes()
    document.close()
    return data


def test_landscape_kk_maps_all_columns_and_skips_empty_slots():
    result = extract_pdf_document(synthetic_kk_pdf())
    header, members = result.bundle.header, result.bundle.members
    assert result.metadata['parser'] == 'kk-landscape-v3'
    assert header.no_kk == '3311111111111111'
    assert header.nama_kepala_keluarga == 'BUDI SANTOSO'
    assert (header.rt, header.rw, header.kode_pos) == ('001', '002', '59511')
    assert (header.desa, header.kecamatan, header.kabupaten, header.provinsi) == ('LAMBANGGELUN', 'KARANGANYAR', 'DEMAK', 'JAWA TENGAH')
    assert len(members) == 2
    assert members[0].nik == '3311111111111112'
    assert members[0].status_hubungan == 'KEPALA KELUARGA'
    assert members[0].status_perkawinan == 'KAWIN TERCATAT'
    assert members[0].nama_ayah == 'SUGENG'
    assert members[0].nama_ibu == 'SRI'
    assert members[0].tanggal_lahir.isoformat() == '1980-01-01'
    assert members[1].nama_lengkap == 'ANI SANTOSO'
    assert validate_extraction(header, members, result.metadata['row_mismatches']) == []
    assert result.thumbnail_data


@pytest.mark.parametrize('edges', [
    [25, 200.8, 282, 324, 425, 467, 518, 640, 783, 841.9],
    [25, 186.2, 276, 325, 420, 480, 530, 640, 780, 841.9],
    [25, 153.5, 270, 330, 420, 480, 530, 640, 780, 841.9],
    [25, 162.4, 268, 325, 410, 472, 525, 635, 778, 841.9],
])
def test_column_anchors_follow_autofit_widths_per_document(edges):
    result = extract_pdf_document(synthetic_kk_pdf(primary_edges=edges))
    assert result.bundle.members[0].nama_lengkap == 'BUDI SANTOSO'
    assert result.bundle.members[0].nik == '3311111111111112'
    assert result.bundle.members[0].tanggal_lahir.isoformat() == '1980-01-01'


def test_vector_grid_uses_real_table_right_instead_of_page_edge():
    primary_edges = [27.3, 186.2, 268, 309.9, 406.1, 455.2, 510.1, 635, 786.4, 829.3]
    secondary_edges = [27.3, 109.1, 160.1, 258.3, 335.7, 401.2, 466.7, 627.2, 829.3]
    result = extract_pdf_document(synthetic_kk_pdf(
        primary_edges=primary_edges, secondary_edges=secondary_edges, draw_grid=True,
    ))
    member = result.bundle.members[0]
    assert member.nama_lengkap == 'BUDI SANTOSO'
    assert member.tanggal_lahir.isoformat() == '1980-01-01'
    assert member.status_hubungan == 'KEPALA KELUARGA'
    assert member.nama_ibu == 'SRI'


def test_secondary_row_anchor_follows_three_line_heading_height():
    result = extract_pdf_document(synthetic_kk_pdf(secondary_anchor_y=303))
    assert len(result.bundle.members) == 2
    assert result.bundle.members[0].status_hubungan == 'KEPALA KELUARGA'


def test_empty_marriage_date_is_valid_for_unrecorded_marriage():
    result = extract_pdf_document(synthetic_kk_pdf(
        first_marital_status='KAWIN BELUM TERCATAT', first_marriage_date=None,
    ))
    member = result.bundle.members[0]
    assert member.status_perkawinan == 'KAWIN BELUM TERCATAT'
    assert member.status_hubungan == 'KEPALA KELUARGA'
    assert validate_extraction(result.bundle.header, result.bundle.members, []) == []


def test_missing_secondary_member_row_is_rejected():
    with pytest.raises(ScannerError) as error:
        extract_pdf_document(synthetic_kk_pdf(missing_secondary_row=True))
    assert error.value.code == 'PDF_TABLE_UNREADABLE'


def test_native_pdf_rejects_pdf_without_selectable_text():
    with pytest.raises(ScannerError) as error:
        extract_pdf_document(synthetic_kk_pdf(include_text=False))
    assert error.value.code == 'PDF_NO_SELECTABLE_TEXT'


def test_native_pdf_rejects_multiple_pages():
    with pytest.raises(ScannerError) as error:
        extract_pdf_document(synthetic_kk_pdf(pages=2))
    assert error.value.code == 'PDF_TOO_MANY_PAGES'


def _scan_session():
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(engine)
    session = Session(engine)
    batch = ScanBatch(batch_code='TEST-001', status='QUEUED')
    item = ScanItem(batch=batch, item_number=1, original_filename='synthetic.pdf')
    session.add(item)
    session.commit()
    return session, item.id


def test_pdf_processing_records_native_attempt():
    session, item_id = _scan_session()
    try:
        result = asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(synthetic_kk_pdf()), filename='synthetic.pdf'), session))
        assert result['status'] == 'EXTRACTED'
        assert len(result['members']) == 2
        attempt = session.scalar(select(ScanAttempt).where(ScanAttempt.scan_item_id == item_id))
        assert attempt.status == 'SUCCESS'
        assert attempt.provider == 'native_pdf'
        assert attempt.model == 'pymupdf-kk-landscape-v3'
    finally:
        session.close()


def test_retry_response_uses_replaced_members():
    session, item_id = _scan_session()
    try:
        first = asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(synthetic_kk_pdf()), filename='synthetic.pdf'), session))
        assert len(first['members']) == 2
        retried = asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(synthetic_kk_pdf(include_second_member=False)), filename='synthetic.pdf'), session))
        assert retried['status'] == 'EXTRACTED'
        assert len(retried['members']) == 1
        assert len(retried['issues']) == 0
    finally:
        session.close()


def test_batch_approval_and_master_export_include_all_approved_rows():
    session, item_id = _scan_session()
    try:
        asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(synthetic_kk_pdf()), filename='synthetic.pdf'), session))
        item = session.get(ScanItem, item_id)
        approved = batches.approve_extracted(item.batch_id, session)
        assert approved['approved_count'] == 1
        assert approved['batch']['approved'] == 1
        assert approved['batch']['extracted'] == 0

        summary = exports.master_summary(session)
        assert summary == {'kk_count': 1, 'row_count': 2}
        response = exports.download_master(session)
        workbook = load_workbook(BytesIO(response.body))
        assert workbook.active.max_row == 3
        assert response.headers['x-export-kk-count'] == '1'
        assert response.headers['x-export-row-count'] == '2'
    finally:
        session.close()


def test_editing_address_recalculates_canonical_dusun():
    session, item_id = _scan_session()
    try:
        asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(synthetic_kk_pdf()), filename='synthetic.pdf'), session))
        updated = scans.update_kk(item_id, KKUpdate(alamat='DK.BOJONGIRENG RT 001'), session)
        assert updated['kk']['dusun'] == 'BOJONGIRENG'
        cleared = scans.update_kk(item_id, KKUpdate(alamat='Jalan tanpa nama dusun'), session)
        assert cleared['kk']['dusun'] is None
    finally:
        session.close()


def test_duplicate_pdf_and_duplicate_kk_are_not_saved_twice():
    session, first_id = _scan_session()
    try:
        source = synthetic_kk_pdf()
        first = asyncio.run(scans._process(first_id, UploadFile(file=BytesIO(source), filename='first.pdf'), session))
        assert first['status'] == 'EXTRACTED'
        first_item = session.get(ScanItem, first_id)

        identical = ScanItem(batch_id=first_item.batch_id, item_number=2, original_filename='identical.pdf')
        session.add(identical); session.commit()
        duplicate_file = asyncio.run(scans._process(identical.id, UploadFile(file=BytesIO(source), filename='identical.pdf'), session))
        assert (duplicate_file['status'], duplicate_file['failure_code']) == ('FAILED', 'DUPLICATE_DOCUMENT')

        same_household = ScanItem(batch_id=first_item.batch_id, item_number=3, original_filename='same-household.pdf')
        session.add(same_household); session.commit()
        duplicate_kk = asyncio.run(scans._process(same_household.id, UploadFile(file=BytesIO(synthetic_kk_pdf(include_second_member=False)), filename='same-household.pdf'), session))
        assert (duplicate_kk['status'], duplicate_kk['failure_code']) == ('FAILED', 'DUPLICATE_HOUSEHOLD')
    finally:
        session.close()


def test_unexpected_parser_error_marks_item_failed(monkeypatch):
    session, item_id = _scan_session()
    try:
        def raise_error(_data):
            raise RuntimeError('synthetic failure')
        monkeypatch.setattr(scans, 'extract_pdf_document', raise_error)
        result = asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(synthetic_kk_pdf()), filename='synthetic.pdf'), session))
        assert result['status'] == 'FAILED'
        assert result['failure_code'] == 'PROCESSING_FAILED'
        attempt = session.scalar(select(ScanAttempt).where(ScanAttempt.scan_item_id == item_id))
        assert attempt.status == 'FAILED'
    finally:
        session.close()


def test_invalid_upload_is_terminal_and_records_preflight_attempt():
    session, item_id = _scan_session()
    try:
        result = asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(b'not a supported document'), filename='invalid.pdf'), session))
        assert (result['status'], result['failure_code']) == ('FAILED', 'UNSUPPORTED_FORMAT')
        attempt = session.scalar(select(ScanAttempt).where(ScanAttempt.scan_item_id == item_id))
        assert (attempt.status, attempt.provider, attempt.model) == ('FAILED', 'validation', 'upload-preflight')
    finally:
        session.close()


def test_process_endpoint_is_idempotent_after_terminal_result():
    session, item_id = _scan_session()
    try:
        first = asyncio.run(scans._process(item_id, UploadFile(file=BytesIO(synthetic_kk_pdf()), filename='synthetic.pdf'), session))
        repeated = asyncio.run(scans.process_item(item_id, UploadFile(file=BytesIO(b'not the original PDF'), filename='synthetic.pdf'), session))
        attempts = session.scalars(select(ScanAttempt).where(ScanAttempt.scan_item_id == item_id)).all()
        assert repeated['status'] == first['status'] == 'EXTRACTED'
        assert len(attempts) == 1
    finally:
        session.close()
